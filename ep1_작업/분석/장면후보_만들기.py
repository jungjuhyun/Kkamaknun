"""판정표.xlsx에서 장면 후보를 세 종류로 묶어 장면후보.xlsx로 낸다.

- 들리는 장면: 맞음·대체로 맞음 줄이 이어지는 곳
- 안 들리는 장면: 틀림·조금 맞음·못 알아들음 줄이 있는 곳 (읽기 줄은 까막눈다운 장면으로 보냄)
- 까막눈다운 장면: 추측·읽기 줄이 있는 곳

고르는 것은 사용자가 한다. 이 스크립트는 판정표에 있는 줄을 시각 순서로 묶기만 한다.
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
SRC = HERE / "판정표.xlsx"
OUT = HERE / "장면후보.xlsx"

GAP = 20.0        # 같은 장면으로 묶는 최대 간격(초): 앞 줄과 이만큼 안 떨어져 있으면 한 장면
MIN_HEAR = 5      # 들리는 장면으로 올릴 최소 판정 줄 수

HEARD = {"맞음", "대체로 맞음"}
MISSED = {"틀림", "조금 맞음", "못 알아들음"}
LISTEN_TYPES = {"따라함", "옮김", "추측", "못알아들음"}


def sec(t):
    m, s = str(t).split(":")
    return int(m) * 60 + float(s)


def fmt(x):
    m, s = divmod(int(round(x)), 60)
    return f"{m}:{s:02d}"


def group(rows):
    """시각 순서의 줄을 GAP 안에 이어지면 한 덩어리로 묶는다."""
    groups, cur = [], []
    for r in rows:
        if cur and r["초"] - cur[-1]["초"] > GAP:
            groups.append(cur)
            cur = []
        cur.append(r)
    if cur:
        groups.append(cur)
    return groups


def hear_groups(rows):
    """맞음·대체로 맞음 줄을 묶되, 사이에 못 들은 줄이 끼면 거기서 끊는다."""
    groups = []
    for chunk in group([r for r in rows if r["유형"] in LISTEN_TYPES and isinstance(r["판정"], str)]):
        cur = []
        for r in chunk:
            if r["판정"] in HEARD:
                cur.append(r)
            else:
                groups.append(cur)
                cur = []
        groups.append(cur)
    return [g for g in groups if len(g) >= MIN_HEAR]


def main():
    df = pd.read_excel(SRC, sheet_name="판정")
    df["초"] = df["시각"].map(sec)
    rows = df.to_dict("records")

    miss = [r for r in rows if r["유형"] in LISTEN_TYPES and r["판정"] in MISSED]
    kkamak = [r for r in rows if r["유형"] in {"추측", "읽기"}]

    scenes = [("들리는 장면", g) for g in hear_groups(rows)]
    scenes += [("안 들리는 장면", g) for g in group(miss)]
    scenes += [("까막눈다운 장면", g) for g in group(kkamak)]

    scenes.sort(key=lambda x: x[1][0]["초"])

    wb = Workbook()
    ws = wb.active
    ws.title = "장면 후보"
    head = ["후보", "종류", "시작", "끝", "길이(초)", "줄 수", "판정",
            "내 말", "애니 대사", "판정표 번호", "고름", "메모"]
    ws.append(head)

    for i, (kind, g) in enumerate(scenes, 1):
        start, end = g[0]["초"], g[-1]["초"]
        verdicts = pd.Series([r["판정"] for r in g if isinstance(r["판정"], str)]).value_counts()
        verdict = ", ".join(f"{k} {v}" for k, v in verdicts.items())
        mine = " / ".join(f"[{r['시각']}] {r['내 말']}" for r in g)
        ani = []
        for r in g:
            a = r["대상 애니 대사"]
            if isinstance(a, str) and a not in ani:
                ani.append(a)
        nums = f"{g[0]['번호']}~{g[-1]['번호']}" if len(g) > 1 else str(g[0]["번호"])
        ws.append([i, kind, fmt(start), fmt(end), round(end - start), len(g), verdict,
                   mine, " / ".join(ani), nums, None, None])

    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    fills = {
        "들리는 장면": PatternFill("solid", fgColor="E2EFDA"),
        "안 들리는 장면": PatternFill("solid", fgColor="FCE4D6"),
        "까막눈다운 장면": PatternFill("solid", fgColor="DDEBF7"),
    }
    yellow = PatternFill("solid", fgColor="FFFF00")
    widths = [6, 14, 7, 7, 8, 6, 22, 70, 60, 11, 8, 30]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = w
    for row in ws.iter_rows():
        for c in row:
            c.font = font
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (7, 8, 9, 12))
    for c in ws[1]:
        c.font = bold
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 2).fill = fills[ws.cell(r, 2).value]
        ws.cell(r, 11).fill = yellow
        ws.cell(r, 12).fill = yellow
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    dv = DataValidation(type="list", formula1='"핵심,쓸 수 있음,안 씀"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"K2:K{ws.max_row}")

    info = wb.create_sheet("안내")
    lines = [
        ("장면 후보 보는 법", None),
        (None, None),
        ("무엇", "판정표.xlsx의 줄을 시각 순서로 묶은 장면 후보 목록이다. 고르는 것은 사용자가 한다."),
        ("들리는 장면", f"맞음·대체로 맞음 줄이 {GAP:.0f}초 안 간격으로 {MIN_HEAR}줄 이상 이어지는 곳. 사이에 못 들은 줄이 끼면 거기서 끊는다 (초록)"),
        ("안 들리는 장면", f"틀림·조금 맞음·못 알아들음 줄. {GAP:.0f}초 안에 이어지면 한 장면 (주황)"),
        ("까막눈다운 장면", f"추측(뜻을 짐작함)·읽기(화면 글자를 읽음) 줄. {GAP:.0f}초 안에 이어지면 한 장면 (파랑)"),
        ("시각", "녹화 시각(원본 영상 기준). 시작·끝은 첫 줄과 마지막 줄의 말 시작 시각이다."),
        ("판정표 번호", "판정표.xlsx 판정 시트의 번호 칸. 자세한 판정 이유는 거기서 본다."),
        ("채울 칸", "노란 칸만. 고름: 핵심 / 쓸 수 있음 / 안 씀 중에서 고른다. 메모는 자유."),
        ("예시", "후보 1을 영상에 꼭 넣고 싶으면 고름에 \"핵심\", 메모에 \"오프닝으로\"처럼 적는다."),
        ("겹침", "같은 시각이 두 종류에 다 들어갈 수 있다. 예: 추측이 틀린 줄은 안 들리는 장면과 까막눈다운 장면 둘 다에 나온다."),
    ]
    for a, b in lines:
        info.append([a, b])
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 100
    for row in info.iter_rows():
        for c in row:
            c.font = font
            c.alignment = Alignment(vertical="top", wrap_text=True)
    info["A1"].font = Font(name="Arial", size=12, bold=True)

    wb.save(OUT)
    counts = pd.Series([k for k, _ in scenes]).value_counts()
    print(f"{OUT.name}: 후보 {len(scenes)}개", dict(counts))


if __name__ == "__main__":
    sys.exit(main())
